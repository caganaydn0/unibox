"""RAG Engine — PostgreSQL pgvector + Ollama Embeddings

Doküman indeksleme (PDF → chunk → embed → pgvector) ve
cosine similarity ile bağlam-zenginleştirilmiş sorgu.
"""
from __future__ import annotations

import collections
import io
import json
import logging
from datetime import datetime, timezone
from pathlib import Path

import httpx
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pypdf import PdfReader
from sqlalchemy import cast, delete, func, literal_column, select
from sqlalchemy.dialects.postgresql import TSQUERY

from app.config import settings
from app.db.models.document_chunk import DocumentChunk
from app.db.models.knowledge_document import KnowledgeDocument, ProcessingStatus
from app.db.session import AsyncSessionLocal
from app.services.chunking import (
    CHUNKER_SURUMU,
    Parca,
    madde_bazli_bol,
    mevzuat_mi,
)

logger = logging.getLogger(__name__)

# ---- Hibrit arama parametreleri ---------------------------------------------
# Bu değerler artık config.py'de yaşıyor (ortamdan ayarlanabilsin diye) ve
# kalibrasyon notları da oraya taşındı. Buradaki adlar yalnızca okunabilirlik
# için; tek doğruluk kaynağı settings.
CANDIDATE_POOL = settings.RAG_CANDIDATE_POOL
RRF_K = settings.RAG_RRF_K
INTENT_BONUS = settings.RAG_INTENT_BONUS
FTS_NORMALIZATION = settings.RAG_FTS_NORMALIZATION
MAX_CHUNKS_PER_DOC = settings.RAG_MAX_CHUNKS_PER_DOC


def _turkish_tsvector(column):
    """to_tsvector('turkish', column) — regconfig literal olarak gömülür.

    Bind parametresi kullanılırsa PostgreSQL to_tsvector(text, text) imzasını
    arar ve fonksiyon bulunamaz hatası verir; bu yüzden literal_column şart.
    """
    return func.to_tsvector(literal_column("'turkish'"), column)


def _turkish_tsquery(soru: str):
    """Kullanıcı sorusunu VEYA bağlantılı tsquery'ye çevirir.

    plainto_tsquery kullanmıyoruz çünkü tüm kelimeleri VE ile bağlıyor:
    "transkriptimi nasıl alırım" sorgusu, bir dokümanda üç kelimenin de
    geçmesini şart koşar ve pratikte hiçbir şey dönmez (ölçüldü: 0 sonuç).

    Bunun yerine soruyu tsvector'e çevirip lexeme'leri ' | ' ile birleştiriyoruz.
    Böylece Postgres'in kendi Türkçe stopword ve stemmer'ı devreye girer,
    kullanıcı metni doğrudan sorguya gömülmediği için enjeksiyon riski de olmaz.

    Her lexeme quote_literal ile tırnaklanıyor. İki sebep:

    1. DOĞRULUK: tsvector_to_array bize zaten NORMALİZE EDİLMİŞ lexeme'i
       veriyor. Tırnaksız bıraktığımızda tsquery ayrıştırıcısı onu yeniden
       yorumluyor ve bileşik lexeme'leri (URL, dosya yolu, e-posta) bölerek
       tsvector'dekiyle eşleşmeyen bir sorguya çevirebiliyor. Tırnaklı hâli
       tsvector'de duranla birebir eşleşir.

    2. SAĞLAMLIK: tsquery'de `! & | ( ) : < '` özel karakter. Bir lexeme
       bunlardan birini içerirse tırnaksız cast "syntax error in tsquery"
       fırlatır. (search() bunu artık yakalıyor ama o zaman da sözcüksel
       aramayı tamamen kaybederdik.)

    Tüm kelimeler stopword ise sonuç boş tsquery olur; bu hata vermez,
    yalnızca hiçbir satırla eşleşmez.
    """
    return cast(
        func.array_to_string(
            func.array(
                select(func.quote_literal(literal_column("l")))
                .select_from(
                    func.unnest(
                        func.tsvector_to_array(
                            func.to_tsvector(literal_column("'turkish'"), soru)
                        )
                    ).alias("l")
                )
                .scalar_subquery()
            ),
            " | ",
        ),
        TSQUERY,
    )


class RagEngine:
    def __init__(self) -> None:
        self._embedding_model = settings.EMBEDDING_MODEL
        self._ollama_url = settings.OLLAMA_BASE_URL
        self._splitter = RecursiveCharacterTextSplitter(
            chunk_size=settings.RAG_CHUNK_SIZE * 5,   # ~500 kelime ≈ 2500 karakter
            chunk_overlap=settings.RAG_CHUNK_OVERLAP * 5,
            length_function=len,
        )

    # ------------------------------------------------------------------ #
    # Embedding
    # ------------------------------------------------------------------ #
    async def embed_text(self, text: str) -> list[float]:
        """Ollama /api/embeddings endpoint'i ile metin → vektör.

        Not: bge-m3 asimetrik önek ("query: " / "passage: ") İSTEMEZ —
        E5 ve bge-v1.5 ailelerinden farklı olarak sorgu ile doküman aynı
        şekilde gömülür. Eklemeyin.
        """
        async with httpx.AsyncClient(base_url=self._ollama_url, timeout=60.0) as client:
            resp = await client.post(
                "/api/embeddings",
                json={"model": self._embedding_model, "prompt": text},
            )
            resp.raise_for_status()
            vektör = resp.json().get("embedding") or []

        # Boyut doğrulaması. Olmadığında iki hata sınıfı, teşhisi zor
        # pgvector mesajlarına dönüşüyordu:
        #  - boş metin -> [] -> "vector must have at least 1 dimension"
        #  - yanlış model (ör. nomic-embed-text hâlâ yüklü) -> 768 boyut ->
        #    "expected 1024 dimensions, not 768"
        if len(vektör) != settings.EMBEDDING_DIMENSIONS:
            raise ValueError(
                f"Embedding boyutu beklenenden farklı: {len(vektör)} != "
                f"{settings.EMBEDDING_DIMENSIONS}. Model '{self._embedding_model}' "
                f"yüklü mü ve EMBEDDING_DIMENSIONS doğru mu? Model değiştiyse "
                f"migration + scripts/reindex_documents.py gerekir."
            )
        return vektör

    # ------------------------------------------------------------------ #
    # Metin çıkarma & parçalama
    # ------------------------------------------------------------------ #
    def extract_text_from_pdf(self, pdf_bytes: bytes) -> str:
        """pypdf ile PDF byte'larından metin çıkar."""
        reader = PdfReader(io.BytesIO(pdf_bytes))
        pages = [page.extract_text() or "" for page in reader.pages]
        return "\n".join(pages)

    @staticmethod
    def decode_text(data: bytes) -> str:
        """Düz metni çöz — UTF-8, olmazsa Windows-1254.

        Türk kurumlarından gelen metin dosyaları sıklıkla cp1254 (Windows
        Türkçe) kodlanmış oluyor. Kör bir utf-8 decode UnicodeDecodeError
        ile patlıyor ve doküman FAILED oluyordu.
        """
        for kodlama in ("utf-8-sig", "utf-8", "cp1254"):
            try:
                return data.decode(kodlama)
            except UnicodeDecodeError:
                continue
        # Son çare: bozuk baytları kaybetmek, dokümanı tamamen kaybetmekten iyi
        return data.decode("utf-8", errors="replace")

    @staticmethod
    def extract_text_from_docx(data: bytes) -> str:
        from docx import Document  # yerel import: yalnızca .docx yüklenince gerekir

        belge = Document(io.BytesIO(data))
        parçalar = [p.text for p in belge.paragraphs]
        # Tablolar da içerik taşıyor — yönetmelik ekleri sıklıkla tablo hâlinde
        for tablo in belge.tables:
            for satır in tablo.rows:
                parçalar.append(" | ".join(h.text.strip() for h in satır.cells))
        return "\n".join(parçalar)

    def extract_text(self, data: bytes, mime_type: str = "", filename: str = "") -> str:
        """Dosya tipine göre doğru çıkarıcıyı seçer.

        Eskiden index_document KOŞULSUZ olarak PdfReader çağırıyordu; oysa
        config.allowed_mime_types TXT, Markdown ve DOCX'i de kabul ediyor.
        Sonuç: admin panelinden yüklenen her metin dosyası FAILED oluyor ve
        yeniden indeksleme de aynı yola gittiği için kurtarılamıyordu.

        MIME'a tek başına güvenmiyoruz: tarayıcılar Markdown'ı tutarsız
        gönderiyor (text/plain, text/markdown, application/octet-stream).
        Uzantı ikinci sinyal olarak kullanılıyor.
        """
        uzantı = Path(filename).suffix.lower()

        if mime_type == "application/pdf" or uzantı == ".pdf":
            return self.extract_text_from_pdf(data)
        if uzantı == ".docx" or "wordprocessingml" in mime_type:
            return self.extract_text_from_docx(data)
        if mime_type.startswith("text/") or uzantı in {".txt", ".md", ".markdown"}:
            return self.decode_text(data)

        # Bilinmeyen tip: PDF imzası varsa PDF say, yoksa metin olarak dene.
        if data[:5] == b"%PDF-":
            return self.extract_text_from_pdf(data)
        return self.decode_text(data)

    def split_text(self, text: str) -> list[str]:
        """Metni chunk'lara böl (yalnızca içerik — meta veri olmadan).

        Geriye dönük uyumluluk için korunuyor. Yeni kod split_with_meta()
        kullanmalı; meta veri olmadan madde numarası/başlık kaybolur.
        """
        return [p.icerik for p in self.split_with_meta(text)]

    def split_with_meta(self, text: str) -> list[Parca]:
        """Metni yapısına göre böler.

        Mevzuat metni (>= 3 madde) madde bazlı, diğer her şey eski sabit
        uzunluk bölücüsüyle işlenir. Mod seçimi otomatik olduğu için
        çağıranın doküman tipini bilmesi gerekmiyor ve rehber/SSS/DOCX
        metinlerinde davranış aynen korunuyor.
        """
        if mevzuat_mi(text):
            parçalar = madde_bazli_bol(text)
            if parçalar:
                return parçalar
            logger.warning("Mevzuat sayıldı ama madde çıkarılamadı; sabit bölmeye düşülüyor.")

        return [
            Parca(icerik=c, meta={"v": CHUNKER_SURUMU, "kind": "plain"})
            for c in self._splitter.split_text(text)
            if c.strip()
        ]

    # ------------------------------------------------------------------ #
    # İndeksleme
    # ------------------------------------------------------------------ #
    async def _indeksleme_basarisiz(self, doc_id: str, exc: Exception) -> None:
        """Dokümanı FAILED işaretler — KENDİ temiz oturumunda.

        Ayrı oturum şart: hata anındaki oturumda bekleyen (commit edilmemiş)
        DELETE olabilir. O oturumda commit çağırmak silmeyi de kalıcı yapardı.
        """
        try:
            async with AsyncSessionLocal() as session:
                doc = await session.get(KnowledgeDocument, doc_id)
                if doc:
                    doc.status = ProcessingStatus.FAILED
                    doc.processing_error = str(exc)[:2000]
                    await session.commit()
        except Exception:
            logger.exception("Doküman FAILED olarak işaretlenemedi: %s", doc_id)

    async def index_document(self, doc_id: str) -> None:
        """Doküman → metin → chunk → embed → pgvector.

        ÜÇ AYRI OTURUM kullanılıyor; bu bilinçli bir tasarım:

        1. Eski hâli tek bir oturumu tüm embedding HTTP çağrıları boyunca açık
           tutuyordu. 47 chunk'lık bir doküman için havuzdaki bir bağlantı
           (pool_size=5) dakikalarca bloke oluyordu.

        2. Daha ciddisi VERİ KAYBIYDI: eski chunk'lar DELETE ediliyor, sonra
           embedding döngüsü başlıyordu. Döngüde bir Ollama timeout'u olursa
           except bloğundaki commit() bekleyen DELETE'i de KALICI yapıyordu.
           Sonuç: daha önce çalışan bir doküman FAILED + sıfır chunk hâline
           geliyor, üstelik chunk_count güncellenmediği için arayüz hâlâ eski
           sayıyı ("20 chunk") gösteriyordu.

        Artık silme ve yazma tek bir transaction'da, tüm embedding'ler
        hazırlandıktan SONRA yapılıyor. Hata hâlinde eski chunk'lara
        dokunulmaz; doküman aramada çalışmaya devam eder.
        """
        # --- 1. Oturum: dokümanı oku, PROCESSING işaretle, oturumu KAPAT ---
        async with AsyncSessionLocal() as session:
            doc = await session.get(KnowledgeDocument, doc_id)
            if not doc:
                logger.error("Doküman bulunamadı: %s", doc_id)
                return
            veri = doc.file_data
            mime = doc.mime_type
            etiket_json = doc.tags_json
            dosya_adı = doc.original_filename
            doc.status = ProcessingStatus.PROCESSING
            await session.commit()

        # --- 2. Oturumsuz: ağır iş (metin çıkarma + embedding) ---
        try:
            metin = self.extract_text(veri, mime, dosya_adı)
            if not metin.strip():
                raise ValueError(
                    "Dosyadan metin çıkarılamadı (boş dosya veya taranmış/OCR'sız PDF)."
                )
            parçalar = self.split_with_meta(metin)
            if not parçalar:
                raise ValueError("Metin parçalanamadı.")

            gömüler = [await self.embed_text(p.icerik) for p in parçalar]
        except Exception as exc:
            await self._indeksleme_basarisiz(doc_id, exc)
            logger.error("İndeksleme hatası (doc=%s): %s", doc_id, exc, exc_info=True)
            raise

        # --- 3. Oturum: sil + yaz, TEK transaction ---
        try:
            async with AsyncSessionLocal() as session:
                await session.execute(
                    delete(DocumentChunk).where(DocumentChunk.document_id == doc_id)
                )
                session.add_all([
                    DocumentChunk(
                        document_id=doc_id,
                        chunk_index=i,
                        content=parça.icerik,
                        embedding=gömü,
                        tags_json=etiket_json,
                        meta_json=parça.meta,
                        word_count=len(parça.icerik.split()),
                    )
                    for i, (parça, gömü) in enumerate(zip(parçalar, gömüler))
                ])
                doc = await session.get(KnowledgeDocument, doc_id)
                doc.status = ProcessingStatus.INDEXED
                doc.indexed_at = datetime.now(timezone.utc).replace(tzinfo=None)
                doc.chunk_count = len(parçalar)
                doc.processing_error = None
                await session.commit()
        except Exception as exc:
            await self._indeksleme_basarisiz(doc_id, exc)
            logger.error("Chunk yazma hatası (doc=%s): %s", doc_id, exc, exc_info=True)
            raise

        logger.info("Doküman indekslendi: %s (%d chunk)", doc_id, len(parçalar))

        from app.core.ws_manager import ws_manager
        await ws_manager.broadcast_to_admins({
            "type": "document_indexed",
            "doc_id": doc_id,
            "filename": dosya_adı,
            "chunk_count": len(parçalar),
        })

    # ------------------------------------------------------------------ #
    # Sorgu (RAG retrieval)
    # ------------------------------------------------------------------ #
    async def search(
        self, question: str, intent_type: str | None = None, top_k: int | None = None
    ) -> list[tuple[str, str, list[str]]]:
        """Hibrit arama — (içerik, belge etiketi, chunk etiketleri) listesi döner.

        İki bağımsız arama çalıştırılır:
          1. Anlamsal: pgvector cosine benzerliği
          2. Sözcüksel: PostgreSQL 'turkish' tam metin araması (ts_rank)

        Sonuçlar Reciprocal Rank Fusion ile birleştirilir: her chunk için
        skor = Σ 1/(RRF_K + sıra). Böylece iki listede de üstlerde çıkan
        chunk'lar öne geçer; sadece birinde çıkanlar da tamamen elenmez.

        Neden hibrit: saf vektör araması Türkçe'de zayıftı (ölçüm: recall@1
        %19). "Transkript" gibi birebir geçen anahtar kelimeleri sözcüksel
        arama yakalar, anlamsal yakınlığı ise vektör araması taşır.
        """
        return await self.search_many([question], intent_type, top_k)

    async def search_many(
        self,
        sorgular: list[str],
        intent_type: str | None = None,
        top_k: int | None = None,
    ) -> list[tuple[str, str, list[str]]]:
        """Birden fazla sorguyu tek bir sonuç listesinde birleştirir.

        Kullanımı: e-postanın KONUSU ve GÖVDESİ ayrı sorgular olarak
        verilir. Konu satırını gövdeye ekleyip tekrarlayarak ağırlık
        vermektense, iki bağımsız arama çalıştırıp zaten var olan RRF ile
        birleştirmek daha ilkeli — konu kısa ve öz olduğu için kendi başına
        güçlü bir sinyal, gövdeye karıştırıldığında kayboluyor.

        Kapı, füzyon ve çeşitlilik TEK KEZ, birleşik aday havuzuna uygulanır.
        """
        k = top_k or settings.RAG_TOP_K

        # Boş/yalnızca boşluk sorgu: Ollama boş vektör döndürüyor ve pgvector
        # "vector must have at least 1 dimension" ile reddediyor. Sohbet
        # ucunda mesaj uzunluğu doğrulanmadığı için bu erişilebilir bir
        # 500'dü. Aramanın anlamı da yok — erken çık.
        temiz_sorgular = [s.strip() for s in sorgular if s and s.strip()]
        if not temiz_sorgular:
            return []

        görünür = (
            KnowledgeDocument.deleted_at.is_(None),
            KnowledgeDocument.status == ProcessingStatus.INDEXED,
        )

        return self._sadelestir(
            await self._ara_zengin(temiz_sorgular, intent_type, k)
        )

    @staticmethod
    def _sadelestir(zengin: list[tuple]) -> list[tuple[str, str, list[str]]]:
        """Zengin demetleri public (içerik, etiket, etiketler) biçimine indirger."""
        return [(içerik, etiket, tags) for _, _, _, içerik, etiket, tags in zengin]

    async def _ara_zengin(
        self, sorgular: list[str], intent_type: str | None, top_k: int | None
    ) -> list[tuple]:
        """search_many ile aynı iş, ama doc_id/chunk_index'i de taşır.

        Belge sırasına dizme bu bilgiye ihtiyaç duyuyor; public search()
        imzasını kirletmemek için ayrı tutuluyor.
        """
        k = top_k or settings.RAG_TOP_K
        temiz = [s.strip() for s in sorgular if s and s.strip()]
        if not temiz:
            return []

        görünür = (
            KnowledgeDocument.deleted_at.is_(None),
            KnowledgeDocument.status == ProcessingStatus.INDEXED,
        )
        skorlar: dict[str, float] = {}
        en_yakın: float | None = None
        vektör_çalıştı = False
        toplam_fts = 0

        for soru in temiz:
            v_ids, f_ids, mesafe, v_ok = await self._adaylari_getir(soru, görünür)
            vektör_çalıştı = vektör_çalıştı or v_ok
            toplam_fts += len(f_ids)
            if mesafe is not None:
                en_yakın = mesafe if en_yakın is None else min(en_yakın, mesafe)
            for liste in (v_ids, f_ids):
                for sıra, cid in enumerate(liste):
                    skorlar[cid] = skorlar.get(cid, 0.0) + 1.0 / (RRF_K + sıra + 1)

        return await self._siralayip_sec(
            skorlar, en_yakın, vektör_çalıştı, toplam_fts, intent_type, k, temiz[0],
        )

    async def _adaylari_getir(
        self, question: str, görünür: tuple
    ) -> tuple[list[str], list[str], float | None, bool]:
        """Tek sorgu için aday listeleri: (vektör, sözcüksel, en yakın mesafe, vektör çalıştı mı)."""
        # --- 1. Anlamsal arama ---
        #
        # Embedding çağrısı BİLİNÇLİ olarak oturum bloğunun DIŞINDA: eskiden
        # havuzdan bir bağlantı, Ollama gidiş-dönüşü boyunca tutuluyordu
        # (pool_size=5 + overflow=10 → 15 eşzamanlı sohbet havuzu tüketir).
        #
        # İki geri getirim BAĞIMSIZ ARIZA ALANI. Eskiden vektör arama ikisinin
        # de ön koşuluydu: Ollama kapalıysa embed_text istisnası doğrudan
        # sohbet isteğini 500'e çeviriyordu — oysa sözcüksel arama Ollama
        # gerektirmiyor ve tek başına iş görebilir.
        vek_ids: list[str] = []
        en_yakın: float | None = None
        vektör_çalıştı = False
        try:
            question_embedding = await self.embed_text(question)
            async with AsyncSessionLocal() as session:  # noqa: SIM117
                mesafe = DocumentChunk.embedding.cosine_distance(question_embedding)
                vek_stmt = (
                    select(DocumentChunk.id, mesafe.label("mesafe"))
                    .join(KnowledgeDocument, DocumentChunk.document_id == KnowledgeDocument.id)
                    .where(*görünür, DocumentChunk.embedding.isnot(None))
                    # id ile beraberlik kırma: eşit mesafede sıra rastgele
                    # olmasın, ölçümler koşular arasında oynamasın.
                    .order_by(mesafe, DocumentChunk.id)
                    .limit(CANDIDATE_POOL)
                )
                vek_satırlar = (await session.execute(vek_stmt)).all()
            vek_ids = [r[0] for r in vek_satırlar]
            en_yakın = float(vek_satırlar[0][1]) if vek_satırlar else None
            vektör_çalıştı = True
        except Exception:
            logger.warning(
                "Anlamsal arama başarısız (Ollama erişilemiyor olabilir) — "
                "yalnızca sözcüksel arama ile devam ediliyor.", exc_info=True,
            )

        # --- 2. Sözcüksel arama ---
        fts_ids: list[str] = []
        try:
            async with AsyncSessionLocal() as session:
                tsv = _turkish_tsvector(DocumentChunk.content)
                tsq = _turkish_tsquery(question)
                fts_stmt = (
                    select(DocumentChunk.id)
                    .join(KnowledgeDocument, DocumentChunk.document_id == KnowledgeDocument.id)
                    .where(*görünür, tsv.op("@@")(tsq))
                    .order_by(
                        func.ts_rank(tsv, tsq, FTS_NORMALIZATION).desc(),
                        DocumentChunk.id,
                    )
                    .limit(CANDIDATE_POOL)
                )
                fts_ids = list((await session.execute(fts_stmt)).scalars())
        except Exception:
            logger.warning(
                "Sözcüksel arama başarısız — yalnızca anlamsal arama ile "
                "devam ediliyor.", exc_info=True,
            )

        return vek_ids, fts_ids, en_yakın, vektör_çalıştı

    async def _siralayip_sec(
        self,
        skorlar: dict[str, float],
        en_yakın: float | None,
        vektör_çalıştı: bool,
        fts_isabet: int,
        intent_type: str | None,
        k: int,
        günlük_metni: str,
    ) -> list[tuple[str, str, list[str]]]:
        """Birleşik aday havuzuna kapı, intent bonusu ve çeşitlilik uygular."""
        # --- Kapsam kontrolü (füzyondan SONRA, iki sinyale birden bakarak) ---
        #
        # Eskiden bu kontrol vektör aramasının hemen ardındaydı ve tek başına
        # erken return yapıyordu. Sonuç: hibrit aramanın SÖZCÜKSEL yarısı,
        # semantik yarı eşiği geçemediğinde HİÇ çalışmıyordu — üstelik
        # vektör tarafının Türkçe'de zayıf olduğu (recall@1 %19) kodun kendi
        # yorumunda yazıyordu. Kapı bekçisi, zayıf olduğu bilinen bileşendi.
        #
        # Artık VE koşulu: soru ancak HEM anlamsal olarak uzaksa HEM DE
        # hiçbir sözcüksel eşleşme yoksa kapsam dışı sayılır. Koşul bir
        # kesişim olduğu için kapsam içi kayıp matematiksel olarak ancak
        # azalabilir.
        #
        # Not: eşik asıl savunma hattı DEĞİL. Ölçüldü (config.py'deki
        # kalibrasyon notu): eşik tamamen kapalıyken bile prompt kuralı 6
        # kapsam dışı sorunun 6'sını doğru reddetti. Bu yüzden burada aşırı
        # mühendislik yapmıyoruz.
        uzak = en_yakın is None or en_yakın > settings.RAG_MAX_DISTANCE
        if vektör_çalıştı and uzak and not fts_isabet:
            logger.info(
                "Kapsam dışı sayıldı (mesafe=%s > %s, sözcüksel isabet=0): %.80s",
                f"{en_yakın:.3f}" if en_yakın is not None else "yok",
                settings.RAG_MAX_DISTANCE, günlük_metni,
            )
            return []

        if not skorlar:
            return []
        question = günlük_metni

        # --- 4. Adayların içeriğini çek ---
        async with AsyncSessionLocal() as session:
            satırlar = (await session.execute(
                select(
                    DocumentChunk.id,
                    DocumentChunk.content,
                    DocumentChunk.tags_json,
                    DocumentChunk.document_id,
                    DocumentChunk.chunk_index,
                    KnowledgeDocument.original_filename,
                    KnowledgeDocument.description,
                )
                .join(KnowledgeDocument, DocumentChunk.document_id == KnowledgeDocument.id)
                .where(DocumentChunk.id.in_(list(skorlar)))
            )).all()

        # --- 5. Intent bonusu ---
        # Intent zaten tespit edilmiş durumda; o intent'le etiketli chunk'ları
        # öne çekiyoruz. Eskiden burada bir filtre vardı ama etiketsiz chunk'lar
        # her zaman dahil edildiği için pratikte hiçbir etkisi yoktu.
        sonuç = []
        for cid, content, tags_json, doc_id, sıra_no, filename, description in satırlar:
            try:
                etiketler = json.loads(tags_json)
            except Exception:
                etiketler = []
            skor = skorlar[cid]
            if intent_type and intent_type in etiketler:
                skor += INTENT_BONUS
            sonuç.append(
                (cid, skor, content, description or filename, etiketler, doc_id, sıra_no)
            )

        sonuç.sort(key=lambda r: r[1], reverse=True)

        # --- 6. Doküman çeşitliliği + doldurma turu ---
        #
        # İlk tur: tek kaynak tüm slotları kapmasın diye doküman başına sınır.
        #
        # İkinci tur (YENİ): sınır yüzünden k'ya ulaşılamadıysa kalan slotlar
        # sınır yok sayılarak doldurulur. Eskiden bu tur yoktu ve sonuç
        # SESSİZCE k'nın altına düşüyordu: MAX_CHUNKS_PER_DOC=2 ile 7 sonuç
        # için en az 4 ayrı doküman gerekiyordu, konuyla eşleşen 2 doküman
        # varsa en fazla 4 chunk dönüyordu ve hiçbir uyarı yoktu.
        seçilen_id: set[str] = set()
        seçilen: list[tuple] = []   # (skor, doc_id, chunk_index, content, etiket, tags)
        doküman_sayacı: collections.Counter[str] = collections.Counter()

        for cid, skor, content, etiket, tags, doc_id, sıra_no in sonuç:
            if len(seçilen) >= k:
                break
            if doküman_sayacı[doc_id] >= MAX_CHUNKS_PER_DOC:
                continue
            doküman_sayacı[doc_id] += 1
            seçilen_id.add(cid)
            seçilen.append((skor, doc_id, sıra_no, content, etiket, tags))

        if len(seçilen) < k:
            for cid, skor, content, etiket, tags, doc_id, sıra_no in sonuç:
                if len(seçilen) >= k:
                    break
                if cid in seçilen_id:
                    continue
                seçilen_id.add(cid)
                seçilen.append((skor, doc_id, sıra_no, content, etiket, tags))

        if len(seçilen) < k:
            logger.info(
                "Aday havuzu yetersiz: %d/%d chunk döndürüldü (soru: %.60s)",
                len(seçilen), k, question,
            )

        # ALAKA SIRASINDA döner. Belge sırasına dizme burada YAPILMAZ —
        # yalnızca LLM bağlamı kurulurken uygulanır (bkz. _baglam_metni).
        #
        # Bu ayrım ölçümle öğrenildi: dizmeyi buraya koyduğumda recall@1
        # %42'den %24'e düştü. Gerçek bir kalite kaybı değildi; sıralamayı
        # bilinçli olarak bozduğumuz için "hedef ilk sırada mı" sorusu
        # anlamını yitirmişti (@7 hiç değişmedi, küme aynıydı). Aramanın
        # alaka sırasını koruması, sıralama kalitesini ölçebilmenin ön koşulu.
        return seçilen

    @staticmethod
    def _belge_sirasina_diz(seçilen: list[tuple]) -> list[tuple]:
        """Aynı dokümandan gelen chunk'ları BELGE SIRASINA göre dizer.

        Sonuçlar füzyon skoru sırasında geliyor; yani aynı yönetmelikten
        MADDE 34'ten sonra MADDE 12 sunulabiliyor. Mevzuat üzerinde akıl
        yürütmesi istenen bir modele maddeleri ters sırada vermek aktif
        olarak kafa karıştırıcı.

        Dokümanların KENDİ arasındaki sırası korunuyor (en iyi skorlu
        doküman önce), yalnızca doküman İÇİNDE chunk_index'e göre diziliyor.
        """
        en_iyi: dict[str, float] = {}
        for skor, doc_id, *_ in seçilen:
            en_iyi[doc_id] = max(en_iyi.get(doc_id, 0.0), skor)
        return sorted(seçilen, key=lambda r: (-en_iyi[r[1]], r[1], r[2]))

    async def query_spec(self, spec, azami_karakter: int | None = None) -> str:
        """QuerySpec ile arama yapıp bağlam metnini döner.

        spec.sorgular konu ve gövdeyi AYRI sorgular olarak taşır; ikisi
        search_many içinde RRF ile birleştirilir.
        """
        return await self._baglam_kur(spec.sorgular, spec.intent_type, azami_karakter)

    async def query(self, question: str, intent_type: str | None = None) -> str:
        """Hibrit arama sonucunu LLM'e verilecek bağlam metnine dönüştürür."""
        return await self._baglam_kur([question], intent_type, None)

    async def _baglam_kur(
        self, sorgular: list[str], intent_type: str | None, azami_karakter: int | None
    ) -> str:
        """Arama + belge sırasına dizme + bağlam metni.

        Belge sırasına dizme YALNIZCA burada uygulanır; search() alaka
        sırasını korur ki sıralama kalitesi ölçülebilsin.
        """
        zengin = await self._ara_zengin(sorgular, intent_type, None)
        return self._baglam_metni(
            self._sadelestir(self._belge_sirasina_diz(zengin)), azami_karakter
        )

    @staticmethod
    def _baglam_metni(
        parçalar: list[tuple[str, str, list[str]]], azami_karakter: int | None = None
    ) -> str:
        """Chunk listesini LLM bağlam metnine çevirir.

        azami_karakter verilirse bütçeyi aşan chunk'lar TAMAMEN atılır —
        metin ortadan KESİLMEZ. Eskiden çağıran taraf `rag_context[:2500]`
        yapıyordu ve bu, son maddeyi cümle ortasından kesip modele yarım
        hüküm veriyordu. Belge sırasına dizmeden sonra "baştaki chunk en
        alakalı" varsayımı da geçersiz, yani körlemesine kesmek büsbütün
        yanlış hâle geldi.
        """
        if not parçalar:
            return ""

        if azami_karakter is not None:
            bütçeli: list[tuple[str, str, list[str]]] = []
            toplam = 0
            for içerik, etiket, tags in parçalar:
                # ayraç ve başlık payı
                maliyet = len(içerik) + len(etiket) + 20
                if bütçeli and toplam + maliyet > azami_karakter:
                    continue
                bütçeli.append((içerik, etiket, tags))
                toplam += maliyet
            if len(bütçeli) < len(parçalar):
                logger.info(
                    "Bağlam bütçesi: %d/%d chunk kullanıldı (%d karakter sınırı).",
                    len(bütçeli), len(parçalar), azami_karakter,
                )
            parçalar = bütçeli

        context_parts = [
            f"[Kaynak {i} — {etiket}]\n{content}"
            for i, (content, etiket, _) in enumerate(parçalar, 1)
        ]
        return "\n\n---\n\n".join(context_parts)

    # ------------------------------------------------------------------ #
    # Silme
    # ------------------------------------------------------------------ #
    async def delete_document_chunks(self, doc_id: str) -> None:
        """Bir dokümanın tüm chunk'larını PostgreSQL'den sil."""
        async with AsyncSessionLocal() as session:
            result = await session.execute(
                delete(DocumentChunk).where(DocumentChunk.document_id == doc_id)
            )
            await session.commit()
            logger.info("%d chunk silindi (doc=%s).", result.rowcount, doc_id)
